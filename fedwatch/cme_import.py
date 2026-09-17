"""导入 CME 官网 FedWatch 工具导出的历史概率。

在 https://www.cmegroup.com/markets/interest-rates/cme-fedwatch-tool.html 里选中某次会议，
左侧「Historical」框的「Downloads」可以下载这次会议过去一年的每日概率（Excel）。
表头一般是 Date, (0-25), (25-50), ... 每行一个日期、每列一个目标区间（bp）。

CME 没公开文件格式说明，这里按"找到一行里有 3 个以上 (a-b) 样式的表头"来识别，
兼容 xlsx 和另存成 csv 的文件。会议日期优先从文件名里认（例如 FedWatch_20270317.xlsx），
认不出就在表头上方的单元格和 sheet 名里找 FOMC 决议日。
"""
from __future__ import annotations

import calendar
import csv
import datetime as dt
import re
from pathlib import Path

RANGE_RE = re.compile(r"^\(?\s*(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)\s*\)?$")
SUFFIXES = {".xlsx", ".xlsm", ".csv"}


class ImportProblem(ValueError):
    pass


def _to_date(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if v is None:
        return None
    s = str(v).strip()
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%Y/%m/%d", "%m/%d/%y", "%d-%b-%Y", "%d %b %Y",
                "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _range_lower_bp(cell) -> int | None:
    m = RANGE_RE.match(str(cell).strip()) if cell is not None else None
    if not m:
        return None
    lo, hi = float(m.group(1)), float(m.group(2))
    if hi <= lo:
        return None
    if hi <= 20:  # 写成百分数的 3.25-3.50
        lo *= 100
    return int(round(lo))


_MON = {m.lower(): i for i, m in enumerate(calendar.month_abbr) if m}


def _meeting_from_text(text: str, fomc_dates: set[dt.date]) -> dt.date | None:
    """在文件名/sheet 名/标题里找 FOMC 决议日。CME 页面上会议标签写成 "17 Mar27"。"""
    cands = []
    for y, m, d in re.findall(r"(20\d{2})[-_/.]?(\d{2})[-_/.]?(\d{2})", text):
        cands.append((int(y), int(m), int(d)))
    for m, d, y in re.findall(r"(\d{1,2})[-_/.](\d{1,2})[-_/.](20\d{2})", text):
        cands.append((int(y), int(m), int(d)))
    for m, d, y in re.findall(r"(?<!\d)(\d{2})(\d{2})(20\d{2})(?!\d)", text):
        cands.append((int(y), int(m), int(d)))
    # 17 Mar 2027 / 17Mar27 / 17-Mar-2027
    for d, mon, y in re.findall(r"(?<!\d)(\d{1,2})[\s_-]*([A-Za-z]{3})[a-z]*[\s_-]*(\d{4}|\d{2})(?!\d)", text):
        if mon.lower() in _MON:
            cands.append((int(y) if len(y) == 4 else 2000 + int(y), _MON[mon.lower()], int(d)))
    # Mar 17, 2027 / March 17 2027
    for mon, d, y in re.findall(r"([A-Za-z]{3})[a-z]*\.?\s+(\d{1,2}),?\s+(20\d{2})", text):
        if mon.lower() in _MON:
            cands.append((int(y), _MON[mon.lower()], int(d)))
    for y, m, d in cands:
        try:
            day = dt.date(y, m, d)
        except ValueError:
            continue
        if day in fomc_dates:
            return day
    return None


def _read_tables(path: Path) -> list[tuple[str, list[list]]]:
    if path.suffix.lower() == ".csv":
        for enc in ("utf-8-sig", "gbk", "latin-1"):
            try:
                with path.open(encoding=enc, newline="") as f:
                    return [(path.stem, list(csv.reader(f)))]
            except UnicodeDecodeError:
                continue
        raise ImportProblem("CSV 编码认不出来")
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    finally:
        wb.close()


def parse_file(path: Path, fomc_dates: set[dt.date]) -> list[tuple[dt.date, dict[dt.date, dict[int, float]]]]:
    """返回 [(会议日, {观察日: {区间下限bp: 概率0~1}})]"""
    out = []
    for sheet, rows in _read_tables(path):
        header_i, cols = None, {}
        for i, row in enumerate(rows[:60]):
            hits = {j: bp for j, c in enumerate(row) if (bp := _range_lower_bp(c)) is not None}
            if len(hits) >= 3:
                header_i, cols = i, hits
                break
        if header_i is None:
            continue
        above = " ".join(str(c) for r in rows[:header_i] for c in r if c is not None)
        meeting = (_meeting_from_text(path.stem, fomc_dates)
                   or _meeting_from_text(sheet, fomc_dates)
                   or _meeting_from_text(above, fomc_dates))
        if meeting is None:
            # 表头上方有日期单元格（datetime 类型）的情况
            for r in rows[:header_i]:
                for c in r:
                    d = _to_date(c)
                    if d in fomc_dates:
                        meeting = d
            if meeting is None:
                example = next((d for d in sorted(fomc_dates) if d > dt.date.today()), max(fomc_dates))
                raise ImportProblem(
                    f"认不出是哪次会议。请把文件改名带上会议日期，比如 FedWatch_{example:%Y%m%d}{path.suffix}")
        date_col = next((j for j in range(len(rows[header_i])) if j not in cols), 0)
        series: dict[dt.date, dict[int, float]] = {}
        for row in rows[header_i + 1:]:
            if date_col >= len(row):
                continue
            day = _to_date(row[date_col])
            if day is None or day >= meeting:
                continue
            probs = {}
            for j, bp in cols.items():
                if j < len(row) and row[j] not in (None, ""):
                    try:
                        probs[bp] = float(str(row[j]).strip().rstrip("%"))
                    except ValueError:
                        continue
            total = sum(probs.values())
            if total <= 0:
                continue
            if total > 1.5:  # 百分数
                probs = {k: v / 100 for k, v in probs.items()}
            series[day] = {k: v for k, v in probs.items() if v > 0}
        if series:
            out.append((meeting, series))
    if not out:
        raise ImportProblem("没找到 (0-25)、(25-50) 这种区间表头")
    return out


def pending_files(folder: Path, already: dict[str, tuple[int, float]]) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    files = []
    for p in sorted(folder.iterdir()):
        if p.suffix.lower() not in SUFFIXES or p.name.startswith("~$"):
            continue
        st = p.stat()
        if already.get(p.name) == (st.st_size, round(st.st_mtime, 3)):
            continue
        files.append(p)
    return files
